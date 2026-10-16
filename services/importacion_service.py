import csv
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from database import conectar, conexion


COLUMNAS = [
    "Codigo_Recurso", "Nombre_Recurso", "Unidad", "Tipo", "Marca", "Precio",
    "Proveedor_RUC", "Proveedor", "Calidad_1_a_5", "Stock",
    "Plazo_Entrega_Dias", "Fecha_Precio", "Actualizar_Precio", "URL_Fuente",
]

ALIASES = {
    "codigo": "Codigo_Recurso", "codigo_recurso": "Codigo_Recurso", "sku": "Codigo_Recurso",
    "nombre": "Nombre_Recurso", "nombre_recurso": "Nombre_Recurso", "recurso": "Nombre_Recurso",
    "unidad": "Unidad", "tipo": "Tipo", "marca": "Marca", "precio": "Precio",
    "proveedor_ruc": "Proveedor_RUC", "ruc_proveedor": "Proveedor_RUC",
    "proveedor": "Proveedor", "nombre_proveedor": "Proveedor",
    "calidad": "Calidad_1_a_5", "calidad_1_a_5": "Calidad_1_a_5",
    "stock": "Stock", "plazo_entrega_dias": "Plazo_Entrega_Dias", "plazo": "Plazo_Entrega_Dias",
    "fecha": "Fecha_Precio", "fecha_precio": "Fecha_Precio",
    "actualizar": "Actualizar_Precio", "actualizar_precio": "Actualizar_Precio",
    "url": "URL_Fuente", "url_fuente": "URL_Fuente", "fuente": "URL_Fuente",
}

TIPOS = {
    "material": "Material", "materiales": "Material",
    "mano de obra": "Mano de Obra", "mano obra": "Mano de Obra", "personal": "Mano de Obra",
    "equipo": "Equipo", "equipos": "Equipo", "herramienta": "Equipo", "herramientas": "Equipo",
}


def _normalizar_texto(valor):
    texto = unicodedata.normalize("NFKD", str(valor or "").strip().lower())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


def _normalizar_cabecera(valor):
    return _normalizar_texto(valor).replace(" ", "_")


def _numero(valor, requerido=False):
    if valor is None or str(valor).strip() == "":
        if requerido:
            raise ValueError("valor numérico obligatorio")
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("S/", "").replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", ".")
    numero = float(texto)
    if numero < 0:
        raise ValueError("no puede ser negativo")
    return numero


def _fecha(valor):
    if valor in (None, ""):
        return date.today().isoformat()
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            pass
    raise ValueError("fecha inválida; use AAAA-MM-DD o DD/MM/AAAA")


def _leer_archivo(ruta):
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".xlsx":
        libro = load_workbook(ruta, read_only=True, data_only=True)
        hoja = libro["Importacion"] if "Importacion" in libro.sheetnames else libro.active
        filas = list(hoja.iter_rows(values_only=True))
        libro.close()
    elif ruta.suffix.lower() == ".csv":
        texto = ruta.read_text(encoding="utf-8-sig")
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;")
        except csv.Error:
            dialecto = csv.excel
        filas = list(csv.reader(texto.splitlines(), dialect=dialecto))
    else:
        raise ValueError("Formato no admitido. Seleccione un archivo .xlsx o .csv.")
    if not filas:
        raise ValueError("El archivo está vacío.")
    cabeceras = []
    for valor in filas[0]:
        clave = _normalizar_cabecera(valor)
        cabeceras.append(ALIASES.get(clave, valor if valor in COLUMNAS else None))
    requeridas = {"Nombre_Recurso", "Unidad", "Tipo", "Precio"}
    faltantes = requeridas - {c for c in cabeceras if c}
    if faltantes:
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(sorted(faltantes)))
    resultado = []
    for numero, valores in enumerate(filas[1:], start=2):
        if not any(v not in (None, "") for v in valores):
            continue
        resultado.append((numero, {cabeceras[i]: valores[i] for i in range(min(len(cabeceras), len(valores))) if cabeceras[i]}))
    return resultado


def previsualizar_importacion(ruta, empresa_id):
    with conectar() as conn:
        existentes = conn.execute(
            "SELECT id,codigo,nombre,unidad,precio FROM recursos WHERE empresa_id=?", (empresa_id,)
        ).fetchall()
    por_codigo = {str(r[1]).strip().upper(): r for r in existentes if r[1]}
    por_nombre = {(_normalizar_texto(r[2]), _normalizar_texto(r[3])): r for r in existentes}
    nuevas_claves = set()
    filas = []
    for numero, bruto in _leer_archivo(ruta):
        errores = []
        try:
            nombre = str(bruto.get("Nombre_Recurso") or "").strip()
            unidad = str(bruto.get("Unidad") or "").strip().upper()
            codigo = str(bruto.get("Codigo_Recurso") or "").strip().upper()
            tipo_raw = _normalizar_texto(bruto.get("Tipo"))
            tipo = TIPOS.get(tipo_raw)
            precio = _numero(bruto.get("Precio"), requerido=True)
            if not nombre: errores.append("nombre obligatorio")
            if not unidad: errores.append("unidad obligatoria")
            if not tipo: errores.append("tipo no reconocido")
            if precio is not None and precio <= 0: errores.append("precio debe ser mayor a cero")
            ruc = str(bruto.get("Proveedor_RUC") or "").strip()
            proveedor = str(bruto.get("Proveedor") or "").strip()
            if ruc and (not ruc.isdigit() or len(ruc) != 11): errores.append("RUC de proveedor inválido")
            calidad_num = _numero(bruto.get("Calidad_1_a_5"))
            calidad = int(calidad_num) if calidad_num is not None else None
            if calidad is not None and (calidad_num != calidad or not 1 <= calidad <= 5): errores.append("calidad debe ser un entero entre 1 y 5")
            stock = _numero(bruto.get("Stock"))
            plazo_num = _numero(bruto.get("Plazo_Entrega_Dias"))
            plazo = int(plazo_num) if plazo_num is not None else None
            if plazo is not None and plazo_num != plazo: errores.append("plazo debe ser un número entero")
            fecha = _fecha(bruto.get("Fecha_Precio"))
            actualizar = _normalizar_texto(bruto.get("Actualizar_Precio")) in {"si", "s", "yes", "1", "true"}
            recurso = por_codigo.get(codigo) if codigo else None
            recurso = recurso or por_nombre.get((_normalizar_texto(nombre), _normalizar_texto(unidad)))
            clave_nueva = (codigo or "", _normalizar_texto(nombre), _normalizar_texto(unidad))
            if recurso and actualizar:
                accion = "ACTUALIZAR"
            elif recurso and (proveedor or ruc):
                accion = "OFERTA"
            elif recurso:
                accion = "SIN CAMBIO"
            elif clave_nueva in nuevas_claves:
                accion = "OFERTA"
            else:
                accion = "CREAR"
                nuevas_claves.add(clave_nueva)
            filas.append({
                "fila": numero, "valido": not errores, "errores": "; ".join(errores), "accion": accion,
                "recurso_id": recurso[0] if recurso else None, "codigo": codigo, "nombre": nombre,
                "unidad": unidad, "tipo": tipo, "marca": str(bruto.get("Marca") or "").strip(),
                "precio": precio, "proveedor_ruc": ruc, "proveedor": proveedor,
                "calidad": calidad, "stock": stock, "plazo": plazo, "fecha": fecha,
                "actualizar": actualizar, "fuente": str(bruto.get("URL_Fuente") or "").strip(),
            })
        except (ValueError, TypeError) as exc:
            filas.append({"fila": numero, "valido": False, "errores": str(exc), "accion": "ERROR"})
    return filas


def ejecutar_importacion(ruta, empresa_id, usuario_id=None):
    filas = previsualizar_importacion(ruta, empresa_id)
    validas = [f for f in filas if f["valido"]]
    if not validas:
        raise ValueError("No existen filas válidas para importar.")
    nuevos = actualizados = ofertas = 0
    with conexion() as conn:
        existentes = conn.execute(
            "SELECT id,codigo,nombre,unidad FROM recursos WHERE empresa_id=?", (empresa_id,)
        ).fetchall()
        por_codigo = {str(r[1]).strip().upper(): r[0] for r in existentes if r[1]}
        por_nombre = {(_normalizar_texto(r[2]), _normalizar_texto(r[3])): r[0] for r in existentes}
        for fila in validas:
            recurso_id = por_codigo.get(fila["codigo"]) if fila["codigo"] else None
            recurso_id = recurso_id or por_nombre.get((_normalizar_texto(fila["nombre"]), _normalizar_texto(fila["unidad"])))
            if recurso_id is None:
                cursor = conn.execute(
                    """INSERT INTO recursos(nombre,unidad,precio,tipo,empresa_id,codigo,marca)
                       VALUES (?,?,?,?,?,?,?)""",
                    (fila["nombre"], fila["unidad"], fila["precio"], fila["tipo"], empresa_id,
                     fila["codigo"] or None, fila["marca"] or None),
                )
                recurso_id = cursor.lastrowid
                if fila["codigo"]:
                    por_codigo[fila["codigo"]] = recurso_id
                por_nombre[(_normalizar_texto(fila["nombre"]), _normalizar_texto(fila["unidad"]))] = recurso_id
                nuevos += 1
            else:
                conn.execute(
                    """UPDATE recursos SET codigo=COALESCE(NULLIF(?,''),codigo),
                       marca=COALESCE(NULLIF(?,''),marca),tipo=? WHERE id=? AND empresa_id=?""",
                    (fila["codigo"], fila["marca"], fila["tipo"], recurso_id, empresa_id),
                )
                if fila["actualizar"]:
                    anterior = conn.execute("SELECT precio FROM recursos WHERE id=?", (recurso_id,)).fetchone()[0]
                    conn.execute("UPDATE recursos SET precio=? WHERE id=?", (fila["precio"], recurso_id))
                    conn.execute(
                        """INSERT INTO historial_precios
                           (empresa_id,recurso_id,precio_anterior,precio_nuevo,origen)
                           VALUES (?,?,?,?,?)""",
                        (empresa_id, recurso_id, anterior, fila["precio"], "IMPORTACION"),
                    )
                    actualizados += 1
            proveedor_id = None
            if fila["proveedor"] or fila["proveedor_ruc"]:
                if fila["proveedor_ruc"]:
                    encontrado = conn.execute(
                        "SELECT id FROM proveedores WHERE empresa_id=? AND ruc=?", (empresa_id, fila["proveedor_ruc"])
                    ).fetchone()
                else:
                    encontrado = conn.execute(
                        "SELECT id FROM proveedores WHERE empresa_id=? AND nombre=? COLLATE NOCASE", (empresa_id, fila["proveedor"])
                    ).fetchone()
                if encontrado:
                    proveedor_id = encontrado[0]
                else:
                    cursor = conn.execute(
                        "INSERT INTO proveedores(empresa_id,ruc,nombre) VALUES (?,?,?)",
                        (empresa_id, fila["proveedor_ruc"] or None, fila["proveedor"] or f"Proveedor {fila['proveedor_ruc']}"),
                    )
                    proveedor_id = cursor.lastrowid
                conn.execute(
                    """INSERT INTO proveedor_precios
                       (empresa_id,proveedor_id,recurso_id,codigo_proveedor,descripcion,unidad,marca,
                        precio,calidad,stock,plazo_entrega_dias,fuente,fecha_precio)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (empresa_id, proveedor_id, recurso_id, fila["codigo"] or None, fila["nombre"],
                     fila["unidad"], fila["marca"] or None, fila["precio"], fila["calidad"],
                     fila["stock"], fila["plazo"], fila["fuente"] or None, fila["fecha"]),
                )
                ofertas += 1
        omitidos = len(filas) - len(validas)
        conn.execute(
            """INSERT INTO importaciones_recursos
               (empresa_id,usuario_id,archivo,formato,filas,nuevos,actualizados,ofertas,omitidos)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (empresa_id, usuario_id, Path(ruta).name, Path(ruta).suffix.lower().lstrip("."),
             len(filas), nuevos, actualizados, ofertas, omitidos),
        )
        if usuario_id:
            conn.execute(
                "INSERT INTO auditoria(usuario_id,accion,entidad,detalle) VALUES (?,?,?,?)",
                (usuario_id, "RECURSOS_IMPORTADOS", "recursos", f"{Path(ruta).name}: {len(validas)} filas"),
            )
    return {"filas": len(filas), "validas": len(validas), "nuevos": nuevos,
            "actualizados": actualizados, "ofertas": ofertas, "omitidos": len(filas)-len(validas)}
