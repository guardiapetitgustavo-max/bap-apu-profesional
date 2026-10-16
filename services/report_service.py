from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import conectar


AZUL = "1F4E78"


def exportar_reporte_excel(empresa_id, ruta):
    wb=Workbook(); ws=wb.active; ws.title="Resumen"
    with conectar() as conn:
        proyectos=conn.execute(
            """SELECT p.id,p.nombre,p.estado,COALESCE(c.nombre,''),COALESCE(SUM(d.parcial),0)
               FROM proyectos p LEFT JOIN clientes c ON c.id=p.cliente_id
               LEFT JOIN detalle_apu d ON d.proyecto_id=p.id WHERE p.empresa_id=?
               GROUP BY p.id,p.nombre,p.estado,c.nombre ORDER BY p.nombre""",(empresa_id,)
        ).fetchall()
        recursos=conn.execute(
            "SELECT nombre,unidad,precio,tipo FROM recursos WHERE empresa_id=? ORDER BY tipo,nombre",(empresa_id,)
        ).fetchall()
        partidas=conn.execute("SELECT nombre,unidad,costo_referencia FROM partidas ORDER BY nombre").fetchall()
        precios=conn.execute(
            """SELECT r.nombre,COALESCE(p.nombre,''),pp.marca,pp.precio,pp.calidad,pp.stock,
                      pp.plazo_entrega_dias,pp.fecha_precio,COALESCE(pp.fuente,'')
               FROM proveedor_precios pp JOIN recursos r ON r.id=pp.recurso_id
               LEFT JOIN proveedores p ON p.id=pp.proveedor_id
               WHERE pp.empresa_id=? ORDER BY r.nombre,pp.fecha_precio DESC""",(empresa_id,)
        ).fetchall()
        historial=conn.execute(
            """SELECT h.fecha,r.nombre,h.precio_anterior,h.precio_nuevo,h.origen,COALESCE(p.nombre,'')
               FROM historial_precios h JOIN recursos r ON r.id=h.recurso_id
               LEFT JOIN proveedores p ON p.id=h.proveedor_id
               WHERE h.empresa_id=? ORDER BY h.id DESC""",(empresa_id,)
        ).fetchall()
    _hoja(ws,"RESUMEN DE PROYECTOS",["ID","Proyecto","Estado","Cliente","Costo APU (S/)"],proyectos)
    wr=wb.create_sheet("Recursos"); _hoja(wr,"CATÁLOGO DE RECURSOS",["Nombre","Unidad","Precio (S/)","Tipo"],recursos)
    wp=wb.create_sheet("Partidas"); _hoja(wp,"BIBLIOTECA DE PARTIDAS",["Partida","Unidad","Costo referencial (S/)"],partidas)
    wpp=wb.create_sheet("Precios Proveedores"); _hoja(wpp,"PRECIOS DE PROVEEDORES",["Recurso","Proveedor","Marca","Precio (S/)","Calidad","Stock","Entrega (días)","Fecha","Fuente"],precios)
    wh=wb.create_sheet("Historial Precios"); _hoja(wh,"HISTORIAL DE PRECIOS",["Fecha","Recurso","Precio anterior","Precio nuevo","Origen","Proveedor"],historial)
    wb.save(ruta); return ruta


def _hoja(ws,titulo,encabezados,filas):
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(encabezados))
    celda=ws.cell(1,1,titulo); celda.font=Font(size=16,bold=True,color="FFFFFF"); celda.fill=PatternFill("solid",fgColor=AZUL); celda.alignment=Alignment(horizontal="center")
    for col,texto in enumerate(encabezados,1):
        c=ws.cell(3,col,texto); c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=AZUL); c.alignment=Alignment(horizontal="center")
    for fila_idx,fila in enumerate(filas,4):
        for col_idx,valor in enumerate(fila,1): ws.cell(fila_idx,col_idx,valor)
    for col_idx in range(1,len(encabezados)+1):
        largo=max([len(str(ws.cell(r,col_idx).value or "")) for r in range(1,ws.max_row+1)] or [10])
        ws.column_dimensions[get_column_letter(col_idx)].width=min(max(largo+2,12),55)
    ws.freeze_panes="A4"; ws.auto_filter.ref=f"A3:{get_column_letter(len(encabezados))}{ws.max_row}"
