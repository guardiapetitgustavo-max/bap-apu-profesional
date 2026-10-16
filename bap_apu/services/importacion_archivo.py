"""Lectura y normalización de archivos de importación (.xlsx / .csv)."""
import csv
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


COLUMNAS = [
    "Codigo_Recurso", "Nombre_Recurso", "Unidad", "Tipo", "Marca", "Precio",
    "Proveedor_RUC", "Proveedor", "Calidad_1_a_5", "Stock",
    "Plazo_Entrega_Dias", "Fecha_Precio", "Actualizar_Precio", "URL_Fuente",
]

ALIASES = {
    "codigo": "Codigo_Recurso", "codigo_recurso": "Codigo_Recurso", "sku": "Codigo_Recurso",
    "nombre": "Nombre_Recurso", "nombre_recurso": "Nombre_Recurso", "recurso": "Nombre_Recurso",
    "unidad": "Unidad", "tipo": "Tipo", "marca": "Marca", "precio": "Precio",
    "proveedor_ruc": "Proveedor_RUC", "ruc_proveedor": "Proveedor_RUC",
    "proveedor": "Proveedor", "nombre_proveedor": "Proveedor",
    "calidad": "Calidad_1_a_5", "calidad_1_a_5": "Calidad_1_a_5",
    "stock": "Stock", "plazo_entrega_dias": "Plazo_Entrega_Dias", "plazo": "Plazo_Entrega_Dias",
    "fecha": "Fecha_Precio", "fecha_precio": "Fecha_Precio",
    "actualizar": "Actualizar_Precio", "actualizar_precio": "Actualizar_Precio",
    "url": "URL_Fuente", "url_fuente": "URL_Fuente", "fuente": "URL_Fuente",
}

TIPOS = {
    "material": "Material", "materiales": "Material",
    "mano de obra": "Mano de Obra", "mano obra": "Mano de Obra", "personal": "Mano de Obra",
    "equipo": "Equipo", "equipos": "Equipo", "herramienta": "Equipo", "herramientas": "Equipo",
}


def normalizar_texto(valor):
    texto = unicodedata.normalize("NFKD", str(valor or "").strip().lower())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


def normalizar_cabecera(valor):
    return normalizar_texto(valor).replace(" ", "_")


def numero(valor, requerido=False):
    if valor is None or str(valor).strip() == "":
        if requerido:
            raise ValueError("valor numérico obligatorio")
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("S/", "").replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", ".")
    numero = float(texto)
    if numero < 0:
        raise ValueError("no puede ser negativo")
    return numero


def fecha(valor):
    if valor in (None, ""):
        return date.today().isoformat()
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            pass
    raise ValueError("fecha inválida; use AAAA-MM-DD o DD/MM/AAAA")


def leer_archivo(ruta):
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".xlsx":
        libro = load_workbook(ruta, read_only=True, data_only=True)
        hoja = libro["Importacion"] if "Importacion" in libro.sheetnames else libro.active
        filas = list(hoja.iter_rows(values_only=True))
        libro.close()
    elif ruta.suffix.lower() == ".csv":
        texto = ruta.read_text(encoding="utf-8-sig")
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;")
        except csv.Error:
            dialecto = csv.excel
        filas = list(csv.reader(texto.splitlines(), dialect=dialecto))
    else:
        raise ValueError("Formato no admitido. Seleccione un archivo .xlsx o .csv.")
    if not filas:
        raise ValueError("El archivo está vacío.")
    cabeceras = []
    for valor in filas[0]:
        clave = normalizar_cabecera(valor)
        cabeceras.append(ALIASES.get(clave, valor if valor in COLUMNAS else None))
    requeridas = {"Nombre_Recurso", "Unidad", "Tipo", "Precio"}
    faltantes = requeridas - {c for c in cabeceras if c}
    if faltantes:
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(sorted(faltantes)))
    resultado = []
    for numero, valores in enumerate(filas[1:], start=2):
        if not any(v not in (None, "") for v in valores):
            continue
        resultado.append((numero, {cabeceras[i]: valores[i] for i in range(min(len(cabeceras), len(valores))) if cabeceras[i]}))
    return resultado
