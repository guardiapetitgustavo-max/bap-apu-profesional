"""Escritura del libro Excel de reportes."""
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = "1F4E78"


@dataclass(frozen=True)
class HojaReporte:
    nombre_hoja: str
    titulo: str
    encabezados: list[str]
    filas: list[tuple]


def escribir_libro(ruta, hojas: list[HojaReporte]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for hoja in hojas:
        _hoja(wb.create_sheet(hoja.nombre_hoja), hoja.titulo, hoja.encabezados, hoja.filas)
    wb.save(ruta)
    return Path(ruta)


def _hoja(ws, titulo, encabezados, filas):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(encabezados))
    celda = ws.cell(1, 1, titulo)
    celda.font = Font(size=16, bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=AZUL)
    celda.alignment = Alignment(horizontal="center")
    for col, texto in enumerate(encabezados, 1):
        c = ws.cell(3, col, texto)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
    for fila_idx, fila in enumerate(filas, 4):
        for col_idx, valor in enumerate(fila, 1):
            ws.cell(fila_idx, col_idx, valor)
    for col_idx in range(1, len(encabezados) + 1):
        largo = max([len(str(ws.cell(r, col_idx).value or "")) for r in range(1, ws.max_row + 1)] or [10])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(largo + 2, 12), 55)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(encabezados))}{ws.max_row}"
