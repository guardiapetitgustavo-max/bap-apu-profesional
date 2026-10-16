import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bap_apu.arranque import preparar_base
from bap_apu.contexto import Contexto
from bap_apu.data.db import Database
from bap_apu.errores import ErrorValidacion
from bap_apu.services.reportes import ReportesService

FIXTURE_1_4_0 = Path(__file__).parent / "fixtures" / "base_1_4_0_con_datos.db"
HOJAS = ["Resumen", "Recursos", "Partidas", "Precios Proveedores", "Historial Precios"]


def _ctx(db):
    with db.lectura() as conn:
        usuario, nombre, rol = conn.execute(
            """SELECT u.usuario, u.nombre_completo, r.nombre
               FROM usuarios u JOIN roles r ON r.id = u.rol_id WHERE u.id = 1"""
        ).fetchone()
        empresa_id = conn.execute(
            """SELECT empresa_id FROM usuario_empresa WHERE usuario_id = 1
               ORDER BY es_principal DESC, empresa_id LIMIT 1"""
        ).fetchone()[0]
        permisos = frozenset(f[0] for f in conn.execute(
            """SELECT p.codigo FROM rol_permisos rp
               JOIN permisos p ON p.id = rp.permiso_id
               JOIN roles r ON r.id = rp.rol_id WHERE r.nombre = 'Administrador'"""))
    return Contexto(usuario_id=1, usuario=usuario, nombre_completo=nombre, rol=rol,
                    permisos=permisos, empresa_id=empresa_id, token="t")


@pytest.fixture
def migrada(tmp_path):
    copia = tmp_path / "apu.db"
    shutil.copy(FIXTURE_1_4_0, copia)
    base = Database(copia)
    preparar_base(base, tmp_path / "datos")
    return base, _ctx(base)


def test_excel_cinco_hojas(db, ctx_admin, tmp_path):
    ruta = ReportesService(db).exportar_excel(ctx_admin, tmp_path / "reporte.xlsx")
    assert ruta.exists() and ruta.stat().st_size > 4000
    assert load_workbook(ruta).sheetnames == HOJAS


def test_pdf_generado(migrada, tmp_path):
    base, ctx = migrada
    destino = tmp_path / "apu.pdf"
    ReportesService(base).exportar_apu_pdf(ctx, 1, "Proyecto", destino)
    assert destino.read_bytes().startswith(b"%PDF")


def test_pdf_proyecto_vacio(db, ctx_admin, tmp_path):
    with db.transaccion() as conn:
        pid = conn.execute(
            "INSERT INTO proyectos (nombre, empresa_id) VALUES ('Vacio', ?)",
            (ctx_admin.empresa_id,),
        ).lastrowid
    with pytest.raises(ErrorValidacion, match="no tiene recursos asignados"):
        ReportesService(db).exportar_apu_pdf(ctx_admin, pid, "Vacio", tmp_path / "x.pdf")
